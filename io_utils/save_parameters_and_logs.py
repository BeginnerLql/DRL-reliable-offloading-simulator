"""Write task outcomes after both replicas finish, with fixed-schema logs."""

import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

from config.paths import SERVER_INFO_PATH, TRAIN_TASK_PATH, RESULTS_DIR, ensure_dirs
from config.params import params as runtime_params
from core.task import compute_task_utility


def task_latency_components(task_results):
    """Audit first completion and separate winner, loser and all-replica waits."""
    rows = pd.DataFrame(task_results).copy()
    assert rows.server_A_id.notna().all()
    timestamps = [f"replica_{name}_{event}_time" for name in ("A", "B")
                  for event in ("queue_enter", "cpu_start", "finish")]
    # An all-single-replica episode leaves the entire B column as Python None.
    rows[timestamps] = rows[timestamps].astype(float)
    has_b = rows.server_B_id.notna()
    first = ~has_b | (rows.replica_A_finish_time <= rows.replica_B_finish_time)
    winner = {}
    waits = {}
    for name in ("A", "B"):
        active = rows.server_A_id.notna() if name == "A" else has_b
        timestamps = [f"replica_{name}_{event}_time" for event in ("queue_enter", "cpu_start", "finish")]
        assert np.isfinite(rows.loc[active, timestamps].to_numpy(dtype=float)).all()
        assert rows.loc[~active, timestamps].isna().all().all()
        queue, start, finish = (rows[f"replica_{name}_{event}_time"]
                                .astype(float) for event in ("queue_enter", "cpu_start", "finish"))
        assert (queue[active] >= rows.arrival_time[active]).all()
        assert (start[active] >= queue[active]).all() and (finish[active] >= start[active]).all()
        waits[name] = start - queue
    for event in ("queue_enter", "cpu_start", "finish"):
        winner[event] = rows[f"replica_A_{event}_time"].where(first, rows[f"replica_B_{event}_time"])
    rows["winner_upload_s"] = winner["queue_enter"] - rows.arrival_time
    rows["winner_cpu_wait_s"] = winner["cpu_start"] - winner["queue_enter"]
    rows["winner_execution_s"] = winner["finish"] - winner["cpu_start"]
    rows["loser_cpu_wait_s"] = waits["B"].where(first, waits["A"])
    rows["all_replica_cpu_wait_s"] = (waits["A"] + waits["B"].fillna(0)) / (1 + has_b.astype(int))
    rows["winner_server_id"] = pd.to_numeric(rows.server_A_id).astype(float).where(
        first, pd.to_numeric(rows.server_B_id).astype(float))
    rows["latency_decomposition_error_s"] = (
        rows.task_latency - rows.winner_upload_s - rows.winner_cpu_wait_s - rows.winner_execution_s).abs()
    np.testing.assert_allclose(rows.task_completion_time, winner["finish"], rtol=0, atol=1e-9)
    np.testing.assert_allclose(rows.task_latency, winner["finish"] - rows.arrival_time, rtol=0, atol=1e-9)
    assert (rows.latency_decomposition_error_s < 1e-9).all()
    return rows


def summarize_satisfied_reliability(records, keys):
    """Pool actual selected reliabilities over satisfied tasks; empty subsets stay NaN."""
    data = records.assign(_satisfied_reliability=records.pair_reliability.where(records.requirement_satisfied))
    summary = data.groupby(keys).agg(
        satisfied_task_count=("requirement_satisfied", "sum"),
        satisfied_reliability_sum=("_satisfied_reliability", "sum"),
    )
    summary["mean_satisfied_reliability"] = (
        summary.satisfied_reliability_sum / summary.satisfied_task_count.replace(0, np.nan))
    return summary.reset_index()


def summarize_task_results(task_results):
    """Summarize all dispatched tasks and actual replicas, including fallback."""
    outcomes = task_latency_components(task_results)
    outcomes["task_utility"] = compute_task_utility(outcomes.task_qos, outcomes.requirement_satisfied)
    episodes = outcomes.groupby("episode").agg(
        task_count=("task_id", "size"), fallback_count=("is_fallback", "sum"),
        ppo_decision_count=("ppo_decision", "sum"),
        mean_task_latency_s=("task_latency", "mean"), theoretical_rsr=("requirement_satisfied", "mean"),
        mean_qos=("task_utility", "mean"), mean_raw_qos=("task_qos", "mean"),
        utility_sum=("task_utility", "sum"),
    ).reset_index()
    episodes = episodes.merge(summarize_satisfied_reliability(outcomes, ["episode"]),
                              on="episode", validate="one_to_one")
    episodes["fallback_fraction"] = episodes.fallback_count / episodes.task_count
    replica_frames = []
    for name in ("A", "B"):
        frame = outcomes[["episode", "task_id", "computation_demand"]].assign(
            server_id=outcomes[f"server_{name}_id"],
            cpu_wait_s=outcomes[f"replica_{name}_cpu_start_time"] - outcomes[f"replica_{name}_queue_enter_time"],
            cpu_busy_s=outcomes[f"replica_{name}_finish_time"] - outcomes[f"replica_{name}_cpu_start_time"],
            finish_time=outcomes[f"replica_{name}_finish_time"],
        )
        replica_frames.append(frame[frame.server_id.notna()])
    replicas = pd.concat(replica_frames, ignore_index=True)
    episodes = episodes.merge(replicas.groupby("episode").cpu_wait_s.mean().rename("mean_cpu_wait_s"), on="episode")
    components = ["winner_upload_s", "winner_cpu_wait_s", "winner_execution_s",
                  "loser_cpu_wait_s", "all_replica_cpu_wait_s"]
    episodes = episodes.merge(outcomes.groupby("episode")[components].mean().add_prefix("mean_"), on="episode")
    # Historical dual-replica runs have equal per-task and per-replica weights.
    if outcomes.server_B_id.notna().all():
        np.testing.assert_allclose(episodes.mean_cpu_wait_s, episodes.mean_all_replica_cpu_wait_s, rtol=0, atol=1e-12)
    loads = replicas.groupby(["episode", "server_id"]).agg(
        replica_count=("task_id", "size"), execution_cycles=("computation_demand", "sum"),
        cpu_busy_s=("cpu_busy_s", "sum"), mean_cpu_wait_s=("cpu_wait_s", "mean"),
    )
    # Include unselected servers, and measure utilization through both replicas' drain.
    index = pd.MultiIndex.from_product(
        [episodes.episode, range(1, runtime_params.NUM_SERVERS + 1)], names=["episode", "server_id"],
    )
    loads = loads.reindex(index, fill_value=0).reset_index()
    duration = replicas.groupby("episode").finish_time.max().rename("episode_duration_s")
    loads = loads.merge(duration, on="episode")
    loads["cpu_utilization"] = loads.cpu_busy_s / loads.episode_duration_s
    counts = outcomes.groupby(["episode", "safe_pair_count"]).size()
    full_index = pd.MultiIndex.from_product(
        [episodes.episode, range(runtime_params.num_actions + 1)], names=["episode", "safe_pair_count"],
    )
    counts = counts.reindex(full_index, fill_value=0).rename("task_count").reset_index()
    return episodes, counts, loads


def save_params_and_logs(
    params, log_data, task_results, run_id, *, task_parameters_path=TRAIN_TASK_PATH,
    run_metadata=None, server_profiles=None,
):
    ensure_dirs()

    model_name = str(params.model_summary).strip().lower()

    filename = os.path.join(
        RESULTS_DIR, f"{model_name}_{run_id}_{params.EXPERIMENT_TAG}.xlsx",
    )

    server_info = pd.read_excel(SERVER_INFO_PATH)
    # Persist the server profiles used for this run.
    if server_profiles is not None:
        server_info = pd.DataFrame(server_profiles.values()).sort_values("Server_ID")

    task_df = pd.read_excel(task_parameters_path)

    other_model_suffixes = tuple(
        f"_{name}" for name in ("ppo", "dqn", "ddpg") if name != model_name
    )
    params_data = {
        attr: [value]
        for attr, value in vars(params).items()
        if not attr.startswith("_") and not attr.endswith(other_model_suffixes)
    }
    params_data["run_id"] = [run_id]
    params_data.update({key: [value] for key, value in (run_metadata or {}).items()})
    df_params = pd.DataFrame(params_data).transpose().reset_index()
    df_params.columns = ["Parameter", "Value"]

    # Episode reward is total effective QoS, including zero-utility unreliable tasks.
    df_logs = pd.DataFrame(
        [
            {
                "Episode": episode,
                "Episode Reward": episodic_reward,
                "task_Avg_Delay": task_avg_delay,
            }
            for episode, episodic_reward, task_avg_delay in log_data
        ],
        columns=["Episode", "Episode Reward", "task_Avg_Delay"],
    )

    task_result_columns = [
        "episode",
        "task_id",
        "task_size",
        "computation_demand",
        "reliability_requirement",
        "arrival_time",
        "deadline_s", "absolute_deadline_s", "task_qos", "task_utility",
        "action_index", "ppo_action", "safe_mask", "action_mask",
        "is_fallback", "ppo_decision", "ppo_transition_index",
        "ppo_delta_t", "ppo_next_task_id", "safe_pair_count",
        "ppo_reward_owner_task_id", "ppo_interval_reward",
        "server_A_id",
        "server_B_id",
        "replica_A_queue_enter_time",
        "replica_A_cpu_start_time",
        "replica_A_finish_time",
        "replica_B_queue_enter_time",
        "replica_B_cpu_start_time",
        "replica_B_finish_time",
        "task_completion_time",
        "task_latency",
        "replica_A_reliability",
        "replica_B_reliability",
        "pair_reliability",
        "requirement_satisfied",
    ]
    df_task_results = pd.DataFrame(task_results, columns=task_result_columns)

    # Exclusive creation preserves every existing experiment workbook.
    with open(filename, "xb") as output:
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df_params.to_excel(writer, sheet_name="Params", index=False)
            task_df.to_excel(writer, sheet_name="Tasks", index=False)
            server_info.to_excel(writer, sheet_name="Servers", index=False)
            df_logs.to_excel(writer, sheet_name="Logs", index=False)
            df_task_results.to_excel(writer, sheet_name="TaskResults", index=False)
            if "is_fallback" in pd.DataFrame(task_results):
                episodes, counts, loads = summarize_task_results(task_results)
                episodes.to_excel(writer, sheet_name="EpisodeMetrics", index=False)
                counts.to_excel(writer, sheet_name="SafePairCounts", index=False)
                loads.to_excel(writer, sheet_name="ServerLoad", index=False)

    wb = load_workbook(filename)

    ws_logs = wb["Logs"]
    for column, title, axis_title, anchor in (
        (2, "Rewards per Episode", "Reward", "F2"),
        (3, "task_Avg_Delay per Episode", "task_Avg_Delay (s)", "F20"),
    ):
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = axis_title
        chart.x_axis.title = "Episode"
        data = Reference(
            ws_logs, min_col=column, min_row=1,
            max_col=column, max_row=ws_logs.max_row,
        )
        categories = Reference(ws_logs, min_col=1, min_row=2, max_row=ws_logs.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        ws_logs.add_chart(chart, anchor)

    wb.save(filename)
    print("successfully saved logs !")
