"""Safely migrate an existing task workbook to add reliability requirements.

This tool preserves the existing Task_ID, Task_Size, and Computation_Demand
columns.  It does not call the full parameter preprocessor and never reads or
writes server_info.xlsx.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from config.paths import DATA_DIR
from tools.generate_server_and_task_parameters import generate_reliability_requirements

TASK_PARAMETERS_FILENAME = "task_parameters.xlsx"
BASE_COLUMNS = ("Task_ID", "Task_Size", "Computation_Demand")
RELIABILITY_COLUMN = "Reliability_Requirement"
DEFAULT_SEED = 2026


def migrate_task_parameters(
    input_path: str | os.PathLike | None = None,
    *,
    seed: int = DEFAULT_SEED,
) -> pd.DataFrame:
    """Add or deterministically refresh requirements using an atomic replacement."""
    path = (
        Path(input_path)
        if input_path is not None
        else Path(DATA_DIR) / TASK_PARAMETERS_FILENAME
    )
    if not path.exists():
        raise FileNotFoundError(f"Task parameter workbook not found: {path}")

    original = pd.read_excel(path)
    missing_columns = [column for column in BASE_COLUMNS if column not in original.columns]
    if missing_columns:
        raise ValueError(
            "task_parameters.xlsx is missing required columns: "
            + ", ".join(missing_columns)
        )
    unexpected_columns = [
        column
        for column in original.columns
        if column not in (*BASE_COLUMNS, RELIABILITY_COLUMN)
    ]
    if unexpected_columns:
        raise ValueError(
            "Refusing to discard unexpected task columns: "
            + ", ".join(unexpected_columns)
        )
    if original.empty:
        raise ValueError("task_parameters.xlsx must contain at least one task")
    if original["Task_ID"].duplicated().any():
        raise ValueError("Task_ID values must be unique")

    original_task_columns = original.loc[:, BASE_COLUMNS].copy(deep=True)
    requirements = generate_reliability_requirements(len(original), seed=seed)
    migrated = original_task_columns.copy(deep=True)
    migrated[RELIABILITY_COLUMN] = requirements

    pd.testing.assert_frame_equal(
        migrated.loc[:, BASE_COLUMNS],
        original_task_columns,
        check_exact=True,
    )
    expected_columns = [*BASE_COLUMNS, RELIABILITY_COLUMN]
    if list(migrated.columns) != expected_columns:
        raise AssertionError("migrated task columns are not in the expected order")
    if not np.isfinite(migrated[RELIABILITY_COLUMN].to_numpy(dtype=float)).all():
        raise AssertionError("Reliability_Requirement contains non-finite values")

    temporary_path = path.with_name(f".{path.stem}.migration-{os.getpid()}.tmp.xlsx")
    with pd.ExcelFile(path) as excel_file:
        sheet_name = excel_file.sheet_names[0]
    workbook = load_workbook(path)
    worksheet = workbook[sheet_name]
    header_columns = {
        worksheet.cell(row=1, column=column).value: column
        for column in range(1, worksheet.max_column + 1)
    }
    if any(column not in header_columns for column in BASE_COLUMNS):
        workbook.close()
        raise ValueError("Workbook header does not match its task dataframe")
    reliability_column = header_columns.get(RELIABILITY_COLUMN)
    if reliability_column is None:
        reliability_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=reliability_column, value=RELIABILITY_COLUMN)
    for row_number, requirement in enumerate(requirements, start=2):
        worksheet.cell(row=row_number, column=reliability_column, value=float(requirement))

    try:
        workbook.save(temporary_path)
        workbook.close()
        written = pd.read_excel(temporary_path, sheet_name=sheet_name)
        if list(written.columns) != expected_columns:
            raise AssertionError("written task columns are not in the expected order")
        pd.testing.assert_frame_equal(
            written.loc[:, BASE_COLUMNS],
            original_task_columns,
            check_dtype=False,
            check_exact=True,
        )
        os.replace(temporary_path, path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()

    return migrated


def _parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=Path(DATA_DIR) / TASK_PARAMETERS_FILENAME,
        help="Existing task_parameters.xlsx to migrate.",
    )
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED)
    return parser.parse_args()


def main():
    args = _parse_args()
    migrated = migrate_task_parameters(args.input, seed=args.seed)
    print(f"Task count: {len(migrated)}")
    counts = migrated[RELIABILITY_COLUMN].value_counts().sort_index()
    for requirement, count in counts.items():
        print(f"{requirement:g}: {int(count)}")
    print("Task_ID, Reliability_Requirement")
    for row in migrated[["Task_ID", RELIABILITY_COLUMN]].head(20).itertuples(index=False):
        print(f"{row.Task_ID}, {row.Reliability_Requirement:g}")


if __name__ == "__main__":
    main()
